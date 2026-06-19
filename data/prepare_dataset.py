"""
Data preparation script for Qaida pronunciation checker.

Parses all 16 xlsx files, normalizes schemas, validates audio mappings,
and outputs a clean unified dataset.

Usage:
    python -m qaida_project.data.prepare_dataset
"""

import os
import sys
import json
import wave
import csv
import openpyxl

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from qaida_project.config.lesson_config import BASE_DIR, LESSON_CONFIG, LESSON4_NULL_AUDIO_FIX


def find_column_index(headers, col_name):
    """Find column index by name (case-insensitive, whitespace-stripped)."""
    target = col_name.strip().lower()
    for idx, h in enumerate(headers):
        if h.strip().lower() == target:
            return idx
    return None


def normalize_transliteration(text):
    """Clean and normalize transliteration text."""
    if text is None:
        return None
    t = str(text).strip()
    # Replace curly apostrophe with straight one
    t = t.replace("\u2019", "'")
    # Normalize multiple spaces to single space
    t = " ".join(t.split())
    return t


def get_audio_path(lesson_num, audio_id, audio_dir):
    """Build full path to WAV file."""
    dir_path = os.path.join(BASE_DIR, audio_dir)
    wav_file = f"{audio_id}.wav"
    full_path = os.path.join(dir_path, wav_file)
    return full_path


def get_audio_duration(wav_path):
    """Get duration of WAV file in seconds."""
    try:
        with wave.open(wav_path, "r") as wf:
            return wf.getnframes() / wf.getframerate()
    except Exception:
        return None


def parse_lesson(lesson_num, config):
    """Parse a single lesson xlsx and return list of entry dicts."""
    xlsx_path = os.path.join(BASE_DIR, config["xlsx"])

    if not os.path.exists(xlsx_path):
        print(f"  ERROR: xlsx not found: {xlsx_path}")
        return []

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # Get headers
    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

    # Find required columns by name
    word_id_idx = find_column_index(headers, "word_id")
    lesson_id_idx = find_column_index(headers, "lesson_id")
    type_idx = find_column_index(headers, "type")
    audio_id_idx = find_column_index(headers, "audio_id")
    arabic_idx = find_column_index(headers, config["arabic_col"])
    translit_idx = find_column_index(headers, config["translit_col"])

    # Validate required columns exist
    required = {
        "word_id": word_id_idx,
        "lesson_id": lesson_id_idx,
        "type": type_idx,
        "audio_id": audio_id_idx,
        "arabic": arabic_idx,
        "transliteration": translit_idx,
    }
    missing = [name for name, idx in required.items() if idx is None]
    if missing:
        print(f"  ERROR: Missing columns: {missing}")
        print(f"  Available headers: {headers}")
        wb.close()
        return []

    # Find optional extra columns
    extra_indices = {}
    if "extra_cols" in config:
        for unified_name, xlsx_name in config["extra_cols"].items():
            idx = find_column_index(headers, xlsx_name)
            if idx is not None:
                extra_indices[unified_name] = idx

    entries = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        word_id = str(row[word_id_idx]).strip() if row[word_id_idx] else None
        lesson_id = str(row[lesson_id_idx]).strip() if row[lesson_id_idx] else None
        entry_type = str(row[type_idx]).strip() if row[type_idx] else None
        audio_id = str(row[audio_id_idx]).strip() if row[audio_id_idx] and str(row[audio_id_idx]).strip() != "None" else None
        arabic_text = str(row[arabic_idx]).strip() if row[arabic_idx] else None
        transliteration_raw = normalize_transliteration(row[translit_idx])

        # Fix Lesson 4 null audio_ids
        if audio_id is None and word_id in LESSON4_NULL_AUDIO_FIX:
            audio_id = LESSON4_NULL_AUDIO_FIX[word_id]

        # Skip rows with no audio_id (shouldn't happen after fix)
        if audio_id is None:
            print(f"  WARNING: Skipping L{lesson_num} row {row_num}: NULL audio_id (word_id={word_id})")
            continue

        # Build audio path
        audio_path = get_audio_path(lesson_num, audio_id, config["audio_dir"])

        # Get extra columns
        extras = {}
        for unified_name, idx in extra_indices.items():
            val = row[idx] if idx < len(row) else None
            extras[unified_name] = str(val).strip() if val else None

        entry = {
            "word_id": word_id,
            "lesson_id": lesson_id,
            "lesson_num": lesson_num,
            "type": entry_type,
            "audio_id": audio_id,
            "arabic_text": arabic_text,
            "transliteration": transliteration_raw,
            "audio_path": audio_path.replace("\\", "/"),
            "sub_type": extras.get("sub_type"),
            "base_letter": extras.get("base_letter"),
        }
        entries.append(entry)

    wb.close()
    return entries


def validate_entries(entries):
    """Validate all entries: check audio exists, no nulls in critical fields."""
    issues = []
    valid_entries = []

    for entry in entries:
        entry_issues = []

        # Check critical fields not null
        for field in ["word_id", "lesson_id", "audio_id", "arabic_text", "transliteration"]:
            if entry.get(field) is None:
                entry_issues.append(f"NULL {field}")

        # Check audio file exists
        if not os.path.exists(entry["audio_path"]):
            entry_issues.append(f"Audio not found: {entry['audio_path']}")

        # Check audio is readable
        if os.path.exists(entry["audio_path"]):
            dur = get_audio_duration(entry["audio_path"])
            if dur is None:
                entry_issues.append("Audio file corrupted/unreadable")
            else:
                entry["duration_s"] = round(dur, 3)

        if entry_issues:
            issues.append((entry.get("word_id", "?"), entry_issues))
        else:
            valid_entries.append(entry)

    return valid_entries, issues


def print_summary(entries):
    """Print dataset summary statistics."""
    from collections import Counter

    print("\n" + "=" * 60)
    print("DATASET SUMMARY")
    print("=" * 60)

    print(f"\nTotal entries: {len(entries)}")

    # Per-lesson breakdown
    lesson_counts = Counter(e["lesson_num"] for e in entries)
    print("\nPer-lesson breakdown:")
    for lnum in sorted(lesson_counts):
        lesson_entries = [e for e in entries if e["lesson_num"] == lnum]
        durations = [e["duration_s"] for e in lesson_entries if "duration_s" in e]
        total_dur = sum(durations)
        print(
            f"  L{lnum:2d} ({lesson_entries[0]['type']:>25s}): "
            f"{lesson_counts[lnum]:3d} entries, "
            f"{total_dur:6.1f}s total audio"
        )

    # Transliteration stats
    all_translits = [e["transliteration"] for e in entries]
    unique_translits = set(t.lower() for t in all_translits)
    print(f"\nTransliterations:")
    print(f"  Total: {len(all_translits)}")
    print(f"  Unique (case-insensitive): {len(unique_translits)}")

    # Duration stats
    durations = [e["duration_s"] for e in entries if "duration_s" in e]
    print(f"\nAudio duration:")
    print(f"  Min: {min(durations):.2f}s")
    print(f"  Max: {max(durations):.2f}s")
    print(f"  Mean: {sum(durations) / len(durations):.2f}s")
    print(f"  Total: {sum(durations):.1f}s ({sum(durations) / 60:.1f} min)")

    # Files > 5s
    long_files = [e for e in entries if e.get("duration_s", 0) > 5.0]
    if long_files:
        print(f"\n  Long files (>5s): {len(long_files)}")
        for e in long_files:
            print(f"    {e['audio_id']}: {e['duration_s']:.2f}s ({e['type']})")


def main():
    print("=" * 60)
    print("QAIDA DATA PREPARATION")
    print("=" * 60)
    print(f"Base directory: {BASE_DIR}")

    all_entries = []

    for lesson_num in sorted(LESSON_CONFIG.keys()):
        config = LESSON_CONFIG[lesson_num]
        print(f"\nProcessing Lesson {lesson_num}: {config['type_name']}...")
        entries = parse_lesson(lesson_num, config)
        print(f"  Parsed {len(entries)} entries")
        all_entries.extend(entries)

    print(f"\n{'=' * 60}")
    print(f"Total parsed: {len(all_entries)} entries")

    # Validate
    print("\nValidating entries...")
    valid_entries, issues = validate_entries(all_entries)

    if issues:
        print(f"\nVALIDATION ISSUES ({len(issues)}):")
        for word_id, entry_issues in issues:
            print(f"  {word_id}: {entry_issues}")
    else:
        print("  All entries valid!")

    print(f"\nValid entries: {len(valid_entries)} / {len(all_entries)}")

    # Print summary
    print_summary(valid_entries)

    # Save outputs
    output_dir = os.path.join(BASE_DIR, "qaida_project", "data")
    os.makedirs(output_dir, exist_ok=True)

    # Save JSON
    json_path = os.path.join(output_dir, "unified_dataset.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(valid_entries, f, ensure_ascii=False, indent=2)
    print(f"\nSaved JSON: {json_path} ({len(valid_entries)} entries)")

    # Save CSV
    csv_path = os.path.join(output_dir, "unified_dataset.csv")
    fieldnames = [
        "word_id", "lesson_id", "lesson_num", "type", "audio_id",
        "arabic_text", "transliteration", "audio_path",
        "sub_type", "base_letter", "duration_s",
    ]
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for entry in valid_entries:
            writer.writerow({k: entry.get(k) for k in fieldnames})
    print(f"Saved CSV: {csv_path} ({len(valid_entries)} entries)")

    return valid_entries


if __name__ == "__main__":
    main()
